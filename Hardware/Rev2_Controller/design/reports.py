"""RC1 deliverables generated on the host: BOM xlsx, CPL csv, test-point map, notes PDFs, evidence/supply-chain reports."""
import csv, json, os, re, sys, math
from collections import OrderedDict, defaultdict

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from design import rev2_design as D

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RC = os.path.join(ROOT, "OSBAMS_Rev2_RELEASE_CANDIDATE_1")
BOARD = json.load(open(os.environ.get("BOARDDATA", os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "build", "boarddata.json"))))
NOTFAB = "RC1 - REVIEW CANDIDATE - NOT FOR FABRICATION"


def pkg(fp):
    n = fp.split(":")[1]
    m = re.search(r"(0402|0603|0805|1206|1210|SOT-23-6|SOT-23|SOIC-8|MSOP-10|LQFP-64|QFN-20|SMA|SMB|DIP-4|DO-35|DO-201AD|SOD-123)", n)
    return m.group(1) if m else n


def groups():
    g = OrderedDict()
    for ref, c in D.COMPS.items():
        if c["key"] in ("TP", "MH", "FID"):
            continue
        k = (c["key"], str(c["value"]), c["fp"], c.get("dnp", False))
        g.setdefault(k, []).append(ref)
    return g


def natkey(r):
    m = re.match(r"([A-Za-z]+)(\d+)", r)
    return (m.group(1), int(m.group(2))) if m else (r, 0)


def bom_rows():
    rows = []
    for i, (k, refs) in enumerate(groups().items(), 1):
        c = D.COMPS[refs[0]]
        refs = sorted(refs, key=natkey)
        notes = "; ".join(sorted({D.COMPS[r].get("note", "") for r in refs if D.COMPS[r].get("note")}))[:240]
        rows.append(OrderedDict([("Item", i), ("Reference(s)", ", ".join(refs)), ("Qty", len(refs)), ("Value", c["value"]), ("Description", c["desc"]),
                                 ("Manufacturer", c["mfr"]), ("MPN", c["mpn"]), ("Package", pkg(c["fp"])), ("Footprint", c["fp"]),
                                 ("Evidence Status", c["evid"]), ("Lifecycle", c["life"]), ("Primary Supplier", "Mouser / Digi-Key (authorized) - NOT QUERIED"),
                                 ("Supplier SKU", "NOT CHECKED (distributor sites unreachable from the build environment)"), ("Stock Check Date", "NOT CHECKED"),
                                 ("Approved Alternate", c["alt"] + (" (UNVERIFIED alternate)" if c["alt"] else "")), ("PCBWay Source / Consign", c["src"]),
                                 ("DNP", "yes" if c.get("dnp") else "no"), ("Notes", notes)]))
    return rows


OFFBOARD = [
    ("Eaton M22-PV-K02 E-stop (NC block)", "J2", "owned / user", "wired to J2; one NC circuit is sufficient"),
    ("C&K T102SHZQE ARM switch", "J3", "owned / user", "series element in the 12 V coil chain"),
    ("Durakool DG57CM-5021-76-1012-R relay K1", "J4 (coil), power path off-board", "owned (Newark 10190042)", "main contacts never touch this PCB"),
    ("Bourns RSA-20-50 shunt", "J5 Kelvin leads", "owned", "2.5 mOhm, 20 A / 50 mV"),
    ("15 A fuse + holder, XT60, manual disconnect", "power path", "owned/purchased", "off-board, main discharge path"),
    ("Mean Well XDR-75-12", "J1", "owned", "set and verify 12.0 V"),
    ("Raspberry Pi 5 + touchscreen + external 5 V supply", "J8 USB-C", "owned", "not powered from this PCB"),
    ("TC74A5-3.3VAT remote probe PCB (+ VJ0805Y104JXXAT)", "J7", "owned (TC74 x2, 100 nF x2)", "see probe/ project"),
    ("NUCLEO-L476RG (ST-LINK probe, fallback)", "J9 via jumper wires", "owned", "keep intact; do not desolder"),
    ("Adafruit INA228 breakout x spare", "reference only", "owned", "validation / fallback comparison"),
]


def write_bom(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "BOM"
    rows = bom_rows()
    ws.append([NOTFAB]); ws["A1"].font = Font(bold=True, color="C00000")
    ws.append(list(rows[0].keys()))
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="33415C")
    for r in rows:
        ws.append(list(r.values()))
    for i, col in enumerate(ws.iter_cols(min_row=2), 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ev = row[9].value
        row[9].fill = PatternFill("solid", fgColor={"VERIFIED_LOCAL": "C6EFCE", "USER_RELAYED_MANUFACTURER": "FFEB9C", "UNVERIFIED": "FFC7CE"}.get(ev, "FFFFFF"))
    ws.freeze_panes = "C3"
    w2 = wb.create_sheet("PCB features (not parts)")
    w2.append(["Ref", "What", "Notes"])
    for ref, c in D.COMPS.items():
        if c["key"] in ("TP", "MH", "FID"):
            w2.append([ref, c["desc"], c["value"] if c["key"] == "TP" else ""])
    w3 = wb.create_sheet("Off-board hardware")
    w3.append(["Item", "Interface", "Status", "Notes"])
    for r in OFFBOARD:
        w3.append(list(r))
    w4 = wb.create_sheet("Evidence legend")
    for r in (("VERIFIED_LOCAL", "read by the build environment"), ("USER_RELAYED_MANUFACTURER", "manufacturer data supplied by the user; PDF not opened here"),
              ("UNVERIFIED", "assumption / proposal requiring confirmation")):
        w4.append(list(r))
    wb.save(path)
    csv.writer(open(path.replace(".xlsx", ".csv"), "w", newline="")).writerows([list(rows[0].keys())] + [list(r.values()) for r in rows])
    return rows


def write_cpl(path_smt, path_all):
    fps = BOARD["footprints"]
    hdr = ["Designator", "Val", "Package", "Mid X", "Mid Y", "Rotation", "Layer"]
    smt, allp = [], []
    for ref, f in fps.items():
        c = D.COMPS.get(ref)
        if not c or c["key"] in ("TP", "MH", "FID"):
            continue
        row = [ref, c["value"], pkg(c["fp"]), f"{f['x']:.4f}mm", f"{-f['y']:.4f}mm", f"{f['rot'] % 360:.1f}", "Top" if f["side"] == "F" else "Bottom"]
        allp.append(row)
        if f["smd"] and not f["tht"]:
            smt.append(row)
    for p, rows in ((path_smt, smt), (path_all, allp)):
        with open(p, "w", newline="") as fh:
            w = csv.writer(fh); w.writerow(hdr); w.writerows(sorted(rows, key=lambda r: natkey(r[0])))
    return len(smt), len(allp)


def write_tp_map(path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    fps = BOARD["footprints"]
    fig, ax = plt.subplots(figsize=(11.7, 10.5))
    ax.add_patch(plt.Rectangle((0, 0), 100, 90, fill=False, lw=1.5))
    for ref, f in fps.items():
        l, t, r, b = f["bbox"]
        ax.add_patch(plt.Rectangle((l, t), r - l, b - t, fill=False, lw=0.3, ec="#999"))
        if ref.startswith(("U", "J", "Q", "SW")):
            ax.text((l + r) / 2, (t + b) / 2, ref, ha="center", va="center", fontsize=6, color="#666")
    n = 0
    for ref, c in D.COMPS.items():
        if c["key"] == "TP":
            f = fps[ref]; n += 1
            ax.plot(f["x"], f["y"], "o", color="#c00", ms=5)
            ax.annotate(f"{ref} {c['value']}", (f["x"], f["y"]), xytext=(3, -3), textcoords="offset points", fontsize=6, color="#900")
    ax.set_xlim(-2, 102); ax.set_ylim(92, -2); ax.set_aspect("equal"); ax.set_title(f"OSBAMS Rev.2 Controller - test-point map ({n} points)   {NOTFAB}", fontsize=9)
    ax.set_xlabel("x mm (board origin top-left)"); ax.set_ylabel("y mm")
    with PdfPages(path) as pdf:
        pdf.savefig(fig)
        fig2, ax2 = plt.subplots(figsize=(11.7, 8.3)); ax2.axis("off")
        rows = [[ref, c["value"], c["nets"][1], f"({fps[ref]['x']:.1f}, {fps[ref]['y']:.1f})"] for ref, c in D.COMPS.items() if c["key"] == "TP"]
        tb = ax2.table(cellText=rows, colLabels=["Ref", "Label", "Net", "x,y mm"], loc="center", cellLoc="left")
        tb.auto_set_font_size(False); tb.set_fontsize(7); tb.scale(1, 1.15)
        ax2.set_title("Test-point list", fontsize=10)
        pdf.savefig(fig2)
    return n


def md_to_pdf(md_path, pdf_path, title):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=8.5, leading=11)
    cell = ParagraphStyle("c", parent=body, fontSize=7, leading=8.5)

    def esc(s):
        s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
        s = re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", s)
        return s
    story = [Paragraph(esc(title), ss["Title"]), Paragraph(f"<b><font color='#C00000'>{NOTFAB}</font></b>", body), Spacer(1, 6)]
    lines = open(md_path).read().split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            story.append(Paragraph(esc(ln[2:]), ss["Heading1"]))
        elif ln.startswith("## "):
            story.append(Paragraph(esc(ln[3:]), ss["Heading2"]))
        elif ln.startswith("### "):
            story.append(Paragraph(esc(ln[4:]), ss["Heading3"]))
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-| :]+\|$", lines[i]):
                    rows.append([Paragraph(esc(c.strip()), cell) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            n = max(len(r) for r in rows)
            rows = [r + [Paragraph("", cell)] * (n - len(r)) for r in rows]
            t = Table(rows, colWidths=[(180 * mm) / n] * n, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dde3ee")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 5)]
            continue
        elif ln.startswith("- "):
            story.append(Paragraph("• " + esc(ln[2:]), body))
        elif ln.strip():
            story.append(Paragraph(esc(ln), body))
        else:
            story.append(Spacer(1, 3))
        i += 1
    SimpleDocTemplate(pdf_path, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=title).build(story)
