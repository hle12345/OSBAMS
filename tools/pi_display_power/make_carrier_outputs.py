#!/usr/bin/env python3
"""BOM / CPL / assembly drawings for the Pi Power Carrier RevB (KiCad 10 via kenv). No Gerbers. usage: make_carrier_outputs.py <project dir>"""
import os, sys, csv, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parts_carrier import *
from kenv import kicli
import cairosvg
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
D = sys.argv[1]; K = f"{D}/kicad/{NAME}.kicad_pcb"; REL = os.environ.get("REL", "RC1")
S = f"{D}/docs/_svg"; os.makedirs(S, exist_ok=True)
for tag, layers in (("top", "F.Cu,F.SilkS,F.Fab,Edge.Cuts"), ("bottom", "B.Cu,B.SilkS,Edge.Cuts")):
    kicli("pcb", "export", "svg", "--exclude-drawing-sheet", "--page-size-mode", "2", "-l", layers, *(["--mirror"] if tag == "bottom" else []), "-o", f"{S}/{tag}.svg", K)
    cairosvg.svg2pdf(url=f"{S}/{tag}.svg", write_to=f"{D}/docs/Assembly_drawing_{tag}.pdf"); cairosvg.svg2png(url=f"{S}/{tag}.svg", write_to=f"{D}/docs/Assembly_drawing_{tag}.png", output_width=1800, background_color="white")
kicli("pcb", "export", "pos", "--format", "csv", "--units", "mm", "--side", "front", "-o", f"{S}/pos.csv", K)
mt = {p[0]: p[6] for p in PARTS}; rows = []
for r_ in csv.DictReader(open(f"{S}/pos.csv")):
    r = r_["Ref"]
    if r in mt and r != "R3": rows.append([r, f"{float(r_['PosX']):.3f}", f"{BOARD_H + float(r_['PosY']):.3f}", "Top", f"{float(r_['Rot']):.0f}", mt[r]])
rows.sort(key=lambda x: x[0])
os.makedirs(f"{D}/cpl", exist_ok=True)
with open(f"{D}/cpl/{NAME}_{REL}_CPL.csv", "w", newline="") as f: w = csv.writer(f); w.writerow(["Designator", "Mid X", "Mid Y", "Layer", "Rotation", "Mount"]); w.writerows(rows)
shutil.rmtree(S, ignore_errors=True)
wb = Workbook(); ws = wb.active; ws.title = "BOM"
ws.append(["Designator", "Qty", "Value", "Description", "Manufacturer", "MPN", "Footprint (KiCad)", "Mount", "Sourcing / notes", "Verification status"])
special = {"U1": "Verify PCBWay stock; if not sourceable mark CONSIGNED / customer-supplied (do not substitute without approval)", "F1": "Confirm stock for 0407008.WR (reel 3000)", "F2": "Active; DigiKey-stocked per design brief",
           "J2": "Mounted from the UNDERSIDE (soldered from the top); order the exact MPN"}
LIBNOTE = "Footprint: official KiCad library; NOT cross-checked against the manufacturer PDF"
status = {"U1": "Footprint from the Mean Well RSDW40/RDDW40 spec drawing (2022-05-24); pin positions checked by script (mirrored drawing, rotated 270 deg)",
          "F1": "Datasheet-verified (Littelfuse 407, rev 09/14/20); land pattern = datasheet recommended pads", "F2": "Datasheet-verified (Littelfuse 451/453); footprint matches",
          "TVS1": "Electrical data USER-RELAYED (Littelfuse SMBJ15A row: VRWM 15 V, VBR 16.7-18.5 V, VC 24.4 V @ 24.6 A, IR 1 uA); KiCad D_SMB footprint, cathode pad 1",
          "J_IN": LIBNOTE + "; ratings owner-cited", "J_DISP": LIBNOTE + "; ratings owner-cited", "J2": "MPN user-relayed from the Samtec page; dimensions verified from catalog F-226 (8.51 mm body, 2.64 mm tail); pad 1.7 / hole 1.0 mm to be checked against Samtec's recommended hole at first fit",
          "C2": LIBNOTE + "; verify body D8 / pitch 3.5", "C4": LIBNOTE + "; verify body D10 / pitch 5.0", "R3": "Not fitted at assembly; fit the select-on-test value at calibration"}
groups = {}
for p in PARTS: groups.setdefault(p[4] + "|" + p[1] if p[0].startswith("TP") or p[4] == "GRM188R71H104KA93D" else p[0], []).append(p)
for k, g in groups.items():
    p = g[0]; ws.append([", ".join(x[0] for x in g), len(g), p[1], p[2], p[3], p[4], p[5], p[6], (special.get(p[0], "") + " " + p[7]).strip() or "Confirm stock/lifecycle at order", status.get(p[0], "MPN from design selection; stock + land pattern not verified offline")])
ws.append([]); ws.append(["MECHANICAL HARDWARE (off-board, in the enclosure / assembly)"])
ws.append(["K1, K2, S3, S4", 4, "", "M3 female-female hex standoff 20 mm (carrier -> enclosure floor) + M3 screws; K1/K2 are the KEY posts: length >= 19 mm is required (see Mechanical_fit_check.txt)", "generic", "M3 x 20 mm standoff", "", "Hardware", "K1 and K2 are mandatory; Pi standoffs 7.4 mm for L = 20 mm", "Computed (Stack/Mechanical check)"])
ws.append(["M1, M2", 2, "", "M2.5 female-female hex standoff 11 mm (header-end spacers to the Pi mounting holes) + M2.5 screws", "generic", "M2.5 x 11 mm standoff", "", "Hardware", "Standard HAT length; confirm on a physical Pi 5", "Computed"])
ws.append(["Pi standoffs", 4, "", "M2.5 standoff 7.4 mm for the Pi 5 itself (so that L = 20 mm carrier standoffs reach the floor)", "generic", "M2.5 x 7.4 mm (or the nearest length; recompute L)", "", "Hardware", "L = Pi standoff + 1.6 + 11.0 mm", "Computed"])
ws.append([]); ws.append(["OFF-BOARD HARNESS (not on the PCB)"])
ws.append(["J_IN mate", 1, "", "Micro-Fit 3.0 receptacle housing, 2 circuits", "Molex", "43025-0200", "", "Harness", "Mates J_IN 430450200 (keyed)", "Standard mate of the 43045 header; mating drawing not in project files"])
ws.append(["J_DISP mate", 1, "", "Micro-Fit 3.0 receptacle housing, 2 circuits (display feed)", "Molex", "43025-0200", "", "Harness", "Mates J_DISP 430450200 (keyed); display end: MX1.25 2-pin per the Waveshare package cables (user-relayed)", "As above; confirm the display-end housing"])
ws.append(["Terminals", 4, "", "Micro-Fit 3.0 female crimp terminals 18 AWG / 0.75 mm2 tin: 2 (J_IN) + 2 (J_DISP) (+ spares)", "Molex", "43030-0038", "", "Harness", "Crimp per Molex ATS-638280200: conductor crimp 1.00-1.10 mm, strip 2.54-2.92 mm, pull >= 89 N; tool 63828-0200, locator 63828-0275", "VERIFIED for 18 AWG (Molex ATS-638280200, uploaded)"])
ws.append(["Wire", 1, "", "18 AWG UL1061-type, insulation OD <= 1.85 mm: J_IN 2 wires; J_DISP 2 wires <= 200 mm", "owner", "18 AWG UL1061-type (supplier by owner)", "", "Harness", "", "Specified; supplier open"])
ws.append([]); ws.append(["REMOVED vs RC2 (not ordered): J_OUT 430450400, interposer J1 430450400, 2 x 43025-0400 housings, 8 x 43030-0038 terminals, 4 x 18 AWG Pi-branch wires, R2 footprint, the interposer PCB, 2 x M2.5 x 20 key standoffs"])
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for i, wd in enumerate([16, 6, 14, 52, 18, 28, 26, 8, 52, 46], 1): ws.column_dimensions[get_column_letter(i)].width = wd
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"; wb.save(f"{D}/bom/{NAME}_{REL}_BOM.xlsx"); print("carrier outputs done")
