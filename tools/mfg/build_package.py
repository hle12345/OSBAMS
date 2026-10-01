"""Build manufacturing/PCBWay_OSBAMS_Rev2_PCBA from the KiCad design files.

    python3 -m tools.mfg.build_package

Outputs are CANDIDATES generated without KiCad (see tools/mfg/__init__.py). Nothing here
declares the board production-ready; the checklist lists what is unresolved.
"""
import csv, math, os, shutil, zipfile
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Polygon, Rectangle, Circle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfWriter, PdfReader

from .pcb_model import Board, REPO
from .schem_model import Schematic
from . import fab_outputs, review_checks, parts, schem_render, docs_text

OUT = os.path.join(REPO, "manufacturing", "PCBWay_OSBAMS_Rev2_PCBA")
SUB = {k: os.path.join(OUT, k) for k in ("Gerber", "Drill", "BOM", "PickAndPlace", "Assembly", "Documentation", "Testing")}
BOM_COLUMNS = ["Reference", "Quantity", "Value", "Manufacturer", "Manufacturer Part Number", "Package",
               "Description", "PCBWay/JLC availability", "Notes"]
AVAIL = "UNKNOWN - not checked (no network access to PCBWay/LCSC stock from the build environment)"


def clean():
    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    for d in SUB.values():
        os.makedirs(d)


# ───────────────────────────── Gerber / drill ─────────────────────────────
def build_gerbers(b):
    stack, info = fab_outputs.build(b)
    names = fab_outputs.save(stack, SUB["Gerber"], SUB["Drill"], None)
    # drill report
    tools = {}
    for kind, items in (("PTH (plated)", info["pth"]), ("NPTH (non-plated)", info["npth"])):
        for x, y, d in items:
            tools.setdefault((kind, round(d, 3)), []).append((x, y))
    lines = ["OSBAMS Rev.2 controller PCB - drill report (generated from the KiCad design by tools/mfg)", "",
             f"Board: {b.width:g} x {b.height:g} mm, {b.thickness:g} mm FR4, 2 copper layers", "",
             "Drill origin: lower-left corner of the board outline (Gerber/CPL origin), Y up", ""]
    for (kind, d), pts in sorted(tools.items()):
        lines.append(f"{kind:<18} {d:>6.3f} mm   x{len(pts)}")
    lines += ["", f"Total holes: {len(info['pth'])} plated, {len(info['npth'])} non-plated", "",
              "Hole coordinates (mm):"]
    for (kind, d), pts in sorted(tools.items()):
        for x, y in sorted(pts):
            lines.append(f"  {kind:<18} {d:>6.3f}  X{x:8.3f}  Y{y:8.3f}")
    open(os.path.join(SUB["Drill"], "OSBAMS_Rev2_drill_report.txt"), "w").write("\n".join(lines) + "\n")
    # zip (Gerber + drill files + candidate notice)
    zpath = os.path.join(SUB["Gerber"], "OSBAMS_Rev2_Gerber.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for d, f in names:
            z.write(os.path.join(d, f), f)
        z.writestr("README_CANDIDATE_NOT_KICAD_EXPORT.txt", docs_text.GERBER_ZIP_NOTICE)
    return names, info, zpath


def verify_gerbers(b, info):
    """Re-read what was written and compare with the design (independent of the writer)."""
    from gerbonara import LayerStack
    tmp = os.path.join(OUT, "_verify")
    os.makedirs(tmp)
    for d in (SUB["Gerber"], SUB["Drill"]):
        for f in os.listdir(d):
            if f.endswith((".gbr", ".drl")):
                shutil.copy(os.path.join(d, f), tmp)
    ls = LayerStack.open_dir(tmp)
    res = {}
    ((x0, y0), (x1, y1)) = ls.bounding_box("mm")
    ow = b.edge_width
    res["bbox_mm"] = (round(x1 - x0 - ow, 3), round(y1 - y0 - ow, 3))
    res["layers"] = sorted(f"{s} {u}" for s, u in ls.graphic_layers)
    res["drill_pth_holes"] = len(ls.drill_pth.objects)
    res["drill_npth_holes"] = len(ls.drill_npth.objects)
    res["expected_pth"], res["expected_npth"] = len(info["pth"]), len(info["npth"])
    res["outline_objects"] = len(ls["mechanical", "outline"].objects)
    res["empty_layers"] = sorted(f"{s} {u}" for (s, u), l in ls.graphic_layers.items() if not l.objects)
    res["copper_pad_objects_top"] = sum(1 for o in ls["top", "copper"].objects if o.__class__.__name__ in ("Flash", "Region"))
    n_top = sum(1 for f in b.footprints for p in f["pads"] if p["type"] != "np_thru_hole"
                and any(l in p["layers"] for l in ("*.Cu", "F.Cu")))
    res["expected_pads_top_cu"] = n_top
    # re-read hole positions vs design
    got = sorted((round(o.x, 2), round(o.y, 2)) for o in ls.drill_pth.objects)
    exp = sorted((round(x, 2), round(y, 2)) for x, y, d in info["pth"])
    res["pth_positions_match"] = got == exp
    shutil.rmtree(tmp)
    return res


def previews(b):
    """PNG previews rendered from the written Gerbers (visual round-trip check)."""
    import io, cairosvg
    from PIL import Image
    from gerbonara import LayerStack
    tmp = os.path.join(OUT, "_pv")
    os.makedirs(tmp)
    for d in (SUB["Gerber"], SUB["Drill"]):
        for f in os.listdir(d):
            if f.endswith((".gbr", ".drl")):
                shutil.copy(os.path.join(d, f), tmp)
    ls = LayerStack.open_dir(tmp)
    B = ((-1, -1), (b.width + 1, b.height + 1))
    def png(layer, fg):
        svg = str(layer.to_svg(force_bounds=B, fg=fg, bg="none"))
        return Image.open(io.BytesIO(cairosvg.svg2png(bytestring=svg.encode(), output_width=1100))).convert("RGBA")
    def comp(stack, name):
        base = Image.new("RGBA", (1100, 1100), (20, 60, 20, 255))
        for key, col, al in stack:
            im = png(ls[key], col)
            if al < 1:
                im.putalpha(im.split()[3].point(lambda v: int(v * al)))
            base = Image.alpha_composite(base, im)
        base.convert("RGB").save(os.path.join(SUB["Documentation"], name))
    comp([(("bottom", "copper"), "#b87333", .6), (("top", "copper"), "#e0a040", 1), (("top", "silk"), "#ffffff", 1),
          (("mechanical", "outline"), "#ffff00", 1)], "preview_top_from_gerber.png")
    comp([(("top", "copper"), "#e0a040", .5), (("bottom", "copper"), "#b87333", 1),
          (("mechanical", "outline"), "#ffff00", 1)], "preview_bottom_from_gerber.png")
    shutil.rmtree(tmp)


# ───────────────────────────── BOM / CPL ─────────────────────────────
def build_bom(b):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(BOM_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DFE6F0")
    placed = {f["ref"] for f in b.footprints}
    order = lambda p: min(int("".join(ch for ch in r if ch.isdigit())) for r in p["refs"])
    rows = []
    for p in sorted(parts.PARTS, key=lambda p: (p["refs"][0][0], order(p))):
        assert all(r in placed for r in p["refs"]), p["refs"]
        rows.append(p)
        ws.append([", ".join(p["refs"]), len(p["refs"]), p["value"], p["mfr"], p["mpn"], p["package"],
                   p["desc"], AVAIL, p["notes"]])
    covered = {r for p in parts.PARTS for r in p["refs"]} | {r for m in parts.MECH for r in m["refs"]}
    missing = placed - covered
    assert not missing, f"footprints without BOM row: {missing}"
    for m in parts.MECH:
        ws.append([", ".join(m["refs"]), len(m["refs"]), "M3 mounting hole", "-", "-", "MountingHole_3.2mm_M3 (NPTH 3.2 mm)",
                   m["desc"], "n/a", "Mechanical feature only - not a purchased part."])
    for col, w in zip("ABCDEFGHI", (13, 9, 26, 24, 24, 38, 52, 34, 80)):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.value == parts.NOT_SPEC:
                c.fill = PatternFill("solid", fgColor="FDECEC")
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("External_not_assembled")
    ws2.append(["Item", "Manufacturer", "Part number", "Notes"])
    for r in parts.EXTERNAL:
        ws2.append(list(r))
    for col, w in zip("ABCD", (44, 30, 26, 90)):
        ws2.column_dimensions[col].width = w
    ws3 = wb.create_sheet("Legend")
    for line in docs_text.BOM_LEGEND:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 140
    path = os.path.join(SUB["BOM"], "OSBAMS_Rev2_BOM.xlsx")
    wb.save(path)
    return path


def build_cpl(b):
    smd = [f for f in b.footprints if "smd" in f["attrs"]]
    path = os.path.join(SUB["PickAndPlace"], "OSBAMS_Rev2_CPL.csv")
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Designator", "Val", "Package", "Mid X", "Mid Y", "Rotation", "Layer"])
        for f in sorted(smd, key=lambda f: f["ref"]):
            gx, gy = b.g(f["x"], f["y"])
            w.writerow([f["ref"], f["value"].split(" ")[0], f["lib"].split(":")[1], f"{gx:.4f}mm", f"{gy:.4f}mm",
                        f"{(f['rot'] % 360):.0f}", "Top" if f["layer"] == "F.Cu" else "Bottom"])
    allp = os.path.join(SUB["PickAndPlace"], "OSBAMS_Rev2_CPL_ALL_PARTS_reference_THT_included.csv")
    with open(allp, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Designator", "Mount", "Package", "Mid X", "Mid Y", "Rotation", "Layer"])
        for f in sorted(b.board_parts(), key=lambda f: f["ref"]):
            gx, gy = b.g(f["x"], f["y"])
            w.writerow([f["ref"], "SMT" if "smd" in f["attrs"] else "THT", f["lib"].split(":")[1], f"{gx:.4f}mm", f"{gy:.4f}mm",
                        f"{(f['rot'] % 360):.0f}", "Top" if f["layer"] == "F.Cu" else "Bottom"])
    open(os.path.join(SUB["PickAndPlace"], "README_CPL.txt"), "w").write(docs_text.CPL_README)
    return path, allp


def verify_cpl(b, path):
    rows = list(csv.DictReader(open(path)))
    smd = {f["ref"] for f in b.footprints if "smd" in f["attrs"]}
    refs = [r["Designator"] for r in rows]
    return dict(smd_parts=sorted(smd), cpl_rows=len(rows), all_placed=smd == set(refs),
                duplicates=len(refs) != len(set(refs)), missing_rotation=[r["Designator"] for r in rows if r["Rotation"] == ""],
                layers=sorted({r["Layer"] for r in rows}))


# ───────────────────────────── drawings ─────────────────────────────
def pad_poly(b, p, mirror=False):
    gx, gy = b.g(p["x"], p["y"])
    w, h = p["w"], p["h"]
    if round(p["ang"]) % 180 == 90:
        w, h = h, w
    pts = []
    sh = p["shape"]
    if sh == "circle":
        pts = [(gx + w / 2 * math.cos(t / 20 * 2 * math.pi), gy + w / 2 * math.sin(t / 20 * 2 * math.pi)) for t in range(21)]
    elif sh == "rect":
        pts = [(gx - w / 2, gy - h / 2), (gx + w / 2, gy - h / 2), (gx + w / 2, gy + h / 2), (gx - w / 2, gy + h / 2)]
    else:   # oval / roundrect -> rounded rectangle
        r = min(w, h) / 2 if sh == "oval" else p["rratio"] * min(w, h)
        for cx, cy, a0 in ((w / 2 - r, h / 2 - r, 0), (-w / 2 + r, h / 2 - r, 90), (-w / 2 + r, -h / 2 + r, 180), (w / 2 - r, -h / 2 + r, 270)):
            for k in range(9):
                t = math.radians(a0 + 90 * k / 8)
                pts.append((gx + cx + r * math.cos(t), gy + cy + r * math.sin(t)))
    if mirror:
        pts = [(b.width - x, y) for x, y in pts]
    return pts


GROUPS = [
    ("Controller interface (NUCLEO-L476RG carries the STM32L476RG - NOT on this PCB)", ["J5"], "#d62728"),
    ("INA228 module header (INA228 is on an external breakout)", ["J4"], "#1f77b4"),
    ("Temperature interface (TC74)", ["U1", "C2"], "#2ca02c"),
    ("Relay (contactor-coil) driver", ["Q1", "R1", "R2", "D2", "J2"], "#ff7f0e"),
    ("Protection (reverse-polarity, TVS, bulk cap)", ["D1", "D3", "C1"], "#9467bd"),
    ("Power / E-stop connectors", ["J1", "J3"], "#8c564b"),
    ("UART / Pi interface", ["J6", "C3"], "#17becf"),
]


def _group_bbox(b, refs, margin=2.2):
    xs, ys = [], []
    for f in b.footprints:
        if f["ref"] in refs:
            for p in f["pads"]:
                gx, gy = b.g(p["x"], p["y"])
                xs += [gx - p["w"] / 2, gx + p["w"] / 2]; ys += [gy - p["h"] / 2, gy + p["h"] / 2]
    return min(xs) - margin, min(ys) - margin, max(xs) + margin, max(ys) + margin


def draw_board(b, ax, side="top", highlight=True):
    mir = side == "bottom"
    X = (lambda x: b.width - x) if mir else (lambda x: x)
    ax.add_patch(Rectangle((0, 0), b.width, b.height, fill=True, fc="#eef3ea", ec="black", lw=1.2))
    if side == "bottom":
        ax.text(b.width / 2, b.height / 2, "BOTTOM: no components\nB.Cu = solid GND pour\n(view mirrored)", ha="center", va="center",
                fontsize=11, color="#555", alpha=.7)
    for f in b.footprints:
        if side == "top":
            for kind, sd, w, data in f["silk"]:
                if sd != "F.SilkS":
                    continue
                if kind in ("line", "poly", "fillpoly"):
                    pts = [b.g(x, y) for x, y in data]
                    ax.add_patch(Polygon(pts, closed=False, fill=(kind == "fillpoly"), fc="#333", ec="#333", lw=max(.4, w * 2.2)))
                elif kind == "circle":
                    cx, cy, r = data
                    gx, gy = b.g(cx, cy)
                    ax.add_patch(Circle((gx, gy), r, fill=False, ec="#333", lw=max(.4, w * 2.2)))
        for p in f["pads"]:
            drill = p["drill"]
            if side == "bottom" and p["type"] == "smd":
                continue
            poly = pad_poly(b, p, mirror=mir)
            fc = "#c9a227" if p["type"] != "np_thru_hole" else "#ffffff"
            if p["type"] != "np_thru_hole":
                ax.add_patch(Polygon(poly, closed=True, fc=fc, ec="#6b5a10", lw=.5))
            if drill:
                gx, gy = b.g(p["x"], p["y"])
                ax.add_patch(Circle((X(gx), gy), drill / 2, fc="white", ec="#444", lw=.5))
            if p["num"] in ("1",) and p["type"] != "np_thru_hole" and side == "top":
                gx, gy = b.g(p["x"], p["y"])
                ax.text(gx, gy, "1", fontsize=5, ha="center", va="center", color="#900", fontweight="bold")
    if side == "top":
        for f in b.footprints:
            for t in f["texts"]:
                if t["kind"] == "reference":
                    gx, gy = b.g(t["x"], t["y"])
                    ax.text(gx, gy, t["text"], fontsize=7, ha="center", va="center", fontweight="bold", color="#0b3d91",
                            bbox=dict(fc="white", ec="none", alpha=.75, pad=.6))
        if highlight:
            for name, refs, col in GROUPS:
                for r in refs:
                    x0, y0, x1, y1 = _group_bbox(b, [r], margin=1.6)
                    ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=True, fc=col, alpha=.16, ec=col, lw=1.4))
        for f in b.footprints:
            if f["ref"] in ("D1", "D2", "D3"):
                p1 = [p for p in f["pads"] if p["num"] == "1"][0]
                gx, gy = b.g(p1["x"], p1["y"])
                ax.annotate("pad 1 = CATHODE\n(footprint)", (gx, gy), xytext=(gx + 4, gy + 3.5), fontsize=5.5, color="#a00000",
                            arrowprops=dict(arrowstyle="->", color="#a00000", lw=.7))
            if f["ref"] == "C1":
                p1 = [p for p in f["pads"] if p["num"] == "1"][0]
                gx, gy = b.g(p1["x"], p1["y"])
                ax.annotate("+ (pad 1)", (gx, gy), xytext=(gx + 3, gy - 6), fontsize=5.5, color="#a00000",
                            arrowprops=dict(arrowstyle="->", color="#a00000", lw=.7))
    ax.set_xlim(-3, b.width + 3); ax.set_ylim(-3, b.height + 3)
    ax.set_aspect("equal"); ax.axis("off")
    ax.text(b.width / 2, -2.2, f"{b.width:g} mm", ha="center", fontsize=6)
    ax.text(-2.4, b.height / 2, f"{b.height:g} mm", rotation=90, va="center", fontsize=6)


def build_assembly_drawing(b):
    tmp_pdf = os.path.join(SUB["Assembly"], "_drawing.pdf")
    with PdfPages(tmp_pdf) as pdf:
        for side in ("top", "bottom"):
            fig = plt.figure(figsize=(11.69, 8.27))
            ax = fig.add_axes([0.02, 0.05, 0.62, 0.85])
            draw_board(b, ax, side)
            fig.text(0.03, 0.95, f"OSBAMS Rev.2 controller PCB - assembly drawing - {side.upper()} view" +
                     (" (mirrored)" if side == "bottom" else ""), fontsize=13, fontweight="bold")
            fig.text(0.03, 0.925, "Derived from 'OSBAMS PCB.kicad_pcb' by tools/mfg - NOT a KiCad plot. Pad '1' markers in red; square pads = pin 1.",
                     fontsize=6.5, color="#a00")
            y = 0.86
            if side == "top":
                fig.text(0.66, y, "Highlighted sections", fontsize=9, fontweight="bold"); y -= 0.03
                for name, refs, col in GROUPS:
                    fig.patches.append(Rectangle((0.66, y - 0.004), 0.012, 0.016, transform=fig.transFigure, fc=col, alpha=.5, ec=col))
                    import textwrap
                    wrapped = textwrap.fill(f"{name}: {', '.join(refs)}", 62)
                    fig.text(0.682, y, wrapped, fontsize=6.4, va="bottom"); y -= 0.017 * (wrapped.count(chr(10)) + 1) + 0.02
                y -= 0.01
                for line in docs_text.DRAWING_CALLOUTS:
                    fig.text(0.66, y, line, fontsize=6.4, va="top", color="#a00000" if line.startswith("!") else "black"); y -= 0.028
            else:
                for line in docs_text.DRAWING_BOTTOM:
                    fig.text(0.66, y, line, fontsize=7, va="top"); y -= 0.045
            pdf.savefig(fig); plt.close(fig)
    notes_pdf = os.path.join(SUB["Assembly"], "_notes_page.pdf")
    docs_text.drawing_notes_pdf(b, notes_pdf)
    w = PdfWriter()
    for f in (tmp_pdf, notes_pdf):
        for pg in PdfReader(f).pages:
            w.add_page(pg)
    out = os.path.join(SUB["Assembly"], "OSBAMS_Rev2_Assembly_Drawing.pdf")
    w.write(out)
    os.remove(tmp_pdf); os.remove(notes_pdf)
    return out


# ───────────────────────────── schematic pdf ─────────────────────────────
def build_schematic_pdf(b, s, review):
    pages = []
    cover = os.path.join(SUB["Documentation"], "_cover.pdf")
    docs_text.schematic_cover_pdf(s, review, cover)
    pages.append(cover)
    figs = os.path.join(SUB["Documentation"], "_figs.pdf")
    with PdfPages(figs) as pdf:
        for title, view, sub in [
            ("Complete schematic (single sheet)", (55, 28, 205, 165), "All sections. Label-based connectivity: nets of the same name are connected."),
            ("Section 1 - 12 V power input and protection (J1, D3, D1, C1)", (60, 33, 202, 64), "Reverse-polarity Schottky D3, TVS D1, bulk capacitor C1. NOTE diode pin 1 = ANODE in this symbol library."),
            ("Section 2 - Relay (contactor-coil) driver and E-stop loop (J3, R1, R2, Q1, D2, J2)", (60, 64, 202, 100), "Low-side MOSFET switch; E-stop loop J3 in series with the coil supply (hardware interruption)."),
            ("Section 3 - MCU interface and UART/Pi header (J5, J6)", (60, 99, 202, 138), "STM32L476RG lives on the external NUCLEO-L476RG; only this header is on the PCB. No USB, no ADC channel."),
            ("Section 4 - Sensors: temperature (U1) and INA228 module header (J4)", (60, 136, 202, 162), "I2C bus: TC74 address 0x4D, INA228 address 0x40."),
        ]:
            fig = schem_render.page(s, "OSBAMS Rev.2 controller - " + title, view, sub)
            pdf.savefig(fig); plt.close(fig)
        fig = docs_text.power_tree_figure()
        pdf.savefig(fig); plt.close(fig)
    pages.append(figs)
    tail = os.path.join(SUB["Documentation"], "_tail.pdf")
    docs_text.schematic_tables_pdf(s, review, tail)
    pages.append(tail)
    w = PdfWriter()
    for f in pages:
        for pg in PdfReader(f).pages:
            w.add_page(pg)
    out = os.path.join(SUB["Documentation"], "OSBAMS_Rev2_Schematic.pdf")
    w.write(out)
    for f in pages:
        os.remove(f)
    return out


def main():
    b = Board()
    s = Schematic()
    clean()
    names, info, zpath = build_gerbers(b)
    gv = verify_gerbers(b, info)
    previews(b)
    bom = build_bom(b)
    cpl, allp = build_cpl(b)
    cv = verify_cpl(b, cpl)
    rv = review_checks.run(b, s)
    ctx = dict(board=b, sch=s, gerber=gv, cpl=cv, review=rv, info=info, names=names)
    build_assembly_drawing(b)
    build_schematic_pdf(b, s, rv)
    docs_text.write_all(ctx, SUB, OUT)
    shutil.copy(os.path.join(os.path.dirname(__file__), "export_with_kicad_cli.sh"), SUB["Documentation"])
    return ctx


if __name__ == "__main__":
    c = main()
    print("gerber verification:", c["gerber"])
    print("cpl verification:", c["cpl"])
