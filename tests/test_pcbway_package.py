"""PCBWay package for the controller PCB: files exist, data is consistent, nothing claims readiness."""
import csv, os, sys, unittest, zipfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)
PKG = os.path.join(ROOT, "manufacturing", "PCBWay_OSBAMS_Rev2_PCBA")

from tools.mfg.pcb_model import Board
from tools.mfg.schem_model import Schematic
from tools.mfg import review_checks, parts

REQUIRED = [
    "Gerber/OSBAMS_Rev2_Gerber.zip", "Drill/OSBAMS_Rev2-PTH.drl", "Drill/OSBAMS_Rev2-NPTH.drl",
    "BOM/OSBAMS_Rev2_BOM.xlsx", "PickAndPlace/OSBAMS_Rev2_CPL.csv",
    "Assembly/OSBAMS_Rev2_Assembly_Drawing.pdf", "Assembly/OSBAMS_Rev2_Assembly_Notes.pdf",
    "Documentation/OSBAMS_Rev2_Schematic.pdf", "Documentation/OSBAMS_Rev2_PCB_REVIEW.md",
    "Documentation/OSBAMS_Rev2_Test_Point_Map.pdf", "Documentation/OSBAMS_Rev2_PCBWAY_CHECKLIST.md",
    "Documentation/export_with_kicad_cli.sh", "Testing/OSBAMS_Rev2_Functional_Test.pdf", "README.md"]


class TestPackage(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.b, cls.s = Board(), Schematic()

    def test_all_deliverables_exist(self):
        for r in REQUIRED:
            self.assertTrue(os.path.getsize(os.path.join(PKG, r)) > 0, r)
        for d in ("Gerber", "Drill", "BOM", "PickAndPlace", "Assembly", "Documentation", "Testing"):
            self.assertTrue(os.path.isdir(os.path.join(PKG, d)))

    def test_gerber_zip_contents(self):
        z = zipfile.ZipFile(os.path.join(PKG, "Gerber", "OSBAMS_Rev2_Gerber.zip"))
        names = z.namelist()
        for want in ("F.Cu", "B.Cu", "F.Mask", "B.Mask", "F.SilkS", "B.SilkS", "F.Paste", "B.Paste", "Edge.Cuts"):
            self.assertTrue(any(want in n and n.endswith(".gbr") for n in names), want)
        self.assertEqual(sum(n.endswith(".drl") for n in names), 2)
        self.assertIn("README_CANDIDATE_NOT_KICAD_EXPORT.txt", names)

    def test_gerber_roundtrip_matches_design(self):
        import shutil, tempfile
        from gerbonara import LayerStack
        tmp = tempfile.mkdtemp()
        zipfile.ZipFile(os.path.join(PKG, "Gerber", "OSBAMS_Rev2_Gerber.zip")).extractall(tmp)
        ls = LayerStack.open_dir(tmp)
        (x0, y0), (x1, y1) = ls.bounding_box("mm")
        self.assertAlmostEqual(x1 - x0 - self.b.edge_width, 80.0, places=2)
        self.assertAlmostEqual(y1 - y0 - self.b.edge_width, 80.0, places=2)
        pth = [(p["x"], p["y"]) for f in self.b.footprints for p in f["pads"] if p["type"] == "thru_hole" and p["drill"]]
        self.assertEqual(len(ls.drill_pth.objects), len(pth))
        self.assertEqual(len(ls.drill_npth.objects), 5)
        got = sorted((round(o.x, 2), round(o.y, 2)) for o in ls.drill_pth.objects)
        exp = sorted((round(self.b.g(x, y)[0], 2), round(self.b.g(x, y)[1], 2)) for x, y in pth)
        self.assertEqual(got, exp)
        shutil.rmtree(tmp)

    def test_bom_columns_and_coverage(self):
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(PKG, "BOM", "OSBAMS_Rev2_BOM.xlsx"))["BOM"]
        hdr = [c.value for c in ws[1]]
        self.assertEqual(hdr, ["Reference", "Quantity", "Value", "Manufacturer", "Manufacturer Part Number", "Package",
                               "Description", "PCBWay/JLC availability", "Notes"])
        refs = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            r = dict(zip(hdr, row))
            self.assertTrue(r["Quantity"] == len([x for x in r["Reference"].split(", ")]))
            for k in hdr:
                self.assertTrue(r[k] not in (None, ""), (r["Reference"], k))
            self.assertNotIn(r["Description"].lower(), ("voltage regulator", "resistor", "capacitor", "diode", "connector"))
            refs += r["Reference"].split(", ")
        self.assertEqual(len(refs), len(set(refs)))
        self.assertEqual(set(refs), {f["ref"] for f in self.b.footprints})
        text = " ".join(str(c) for row in ws.iter_rows(min_row=2, values_only=True) for c in row).lower()
        for external in ("fuse holder", "xt60", "shunt", "disconnect", "6060b"):
            self.assertNotIn(external, text.replace("no xt60", ""), external) if external != "xt60" else None

    def test_unspecified_parts_are_flagged_not_invented(self):
        spec = {r for p in parts.PARTS if p["mpn"] == parts.NOT_SPEC for r in p["refs"]}
        self.assertEqual(spec, {"C2", "C3", "R1", "R2", "J1", "J2", "J3", "J4", "J5", "J6"})
        chk = open(os.path.join(PKG, "Documentation", "OSBAMS_Rev2_PCBWAY_CHECKLIST.md")).read()
        self.assertIn("missing MPN: C2, C3, R1, R2, J1, J2, J3, J4, J5, J6", chk)

    def test_cpl_complete(self):
        rows = list(csv.DictReader(open(os.path.join(PKG, "PickAndPlace", "OSBAMS_Rev2_CPL.csv"))))
        smd = {f["ref"] for f in self.b.footprints if "smd" in f["attrs"]}
        self.assertEqual({r["Designator"] for r in rows}, smd)
        self.assertEqual(len(rows), len(smd))
        for r in rows:
            for k in ("Mid X", "Mid Y", "Rotation", "Layer"):
                self.assertNotEqual(r[k], "")
        allp = list(csv.DictReader(open(os.path.join(PKG, "PickAndPlace",
                    "OSBAMS_Rev2_CPL_ALL_PARTS_reference_THT_included.csv"))))
        self.assertEqual({r["Designator"] for r in allp}, {f["ref"] for f in self.b.board_parts()})

    def test_schematic_and_pcb_agree(self):
        rv = review_checks.run(self.b, self.s)
        self.assertEqual(rv["netlist_diffs"], [])
        self.assertEqual(rv["ref_footprint_mismatch"], [])
        self.assertEqual(rv["dangling_track_ends"], [])
        self.assertEqual(rv["courtyard_overlaps"], [])

    def test_diode_polarity_mismatch_is_reproducible(self):
        """Symbol pin 1 = anode (triangle base); KiCad footprint pad 1 = cathode ('K' on silk / band at pad 1)."""
        from tools.mfg.sexp import find
        for lib in ("OSBAMS:D_H",):
            sym = self.s.libs[lib]
            tri = None
            for u in find(sym, "symbol"):
                for e in u[2:]:
                    if isinstance(e, list) and e[0] == "polyline":
                        pts = [(float(p[1]), float(p[2])) for p in find(find(e, "pts")[0], "xy")]
                        if len(pts) == 4:
                            tri = pts
            apex_x = max(p[0] for p in tri); base_x = min(p[0] for p in tri)
            pin1 = [p for u in find(sym, "symbol") for p in find(u, "pin") if find(p, "number")[0][1] == "1"][0]
            self.assertLess(float(find(pin1, "at")[0][1]), base_x)      # pin 1 on the triangle-base side = anode
            self.assertGreater(apex_x, base_x)
        for ref in ("D2", "D3"):
            fp = [f for f in self.b.footprints if f["ref"] == ref][0]
            k = [t for t in fp["texts"] if t["text"] == "K"]
            self.assertTrue(k, f"{ref} footprint has no 'K' marker")
            p1 = [p for p in fp["pads"] if p["num"] == "1"][0]
            p2 = [p for p in fp["pads"] if p["num"] == "2"][0]
            d1 = ((k[0]["x"] - p1["x"]) ** 2 + (k[0]["y"] - p1["y"]) ** 2) ** .5
            d2 = ((k[0]["x"] - p2["x"]) ** 2 + (k[0]["y"] - p2["y"]) ** 2) ** .5
            self.assertLess(d1, d2)                                       # K marker sits at pad 1
        text = open(os.path.join(PKG, "Documentation", "OSBAMS_Rev2_PCB_REVIEW.md")).read()
        self.assertIn("Diode polarity (CRITICAL)", text)

    def test_nothing_claims_production_ready(self):
        for rel in ("Documentation/OSBAMS_Rev2_PCB_REVIEW.md", "Documentation/OSBAMS_Rev2_PCBWAY_CHECKLIST.md", "README.md"):
            t = open(os.path.join(PKG, rel)).read()
            self.assertIn("NOT", t.upper())
            self.assertRegex(t, r"(?i)not (ready for )?production")
        chk = open(os.path.join(PKG, "Documentation", "OSBAMS_Rev2_PCBWAY_CHECKLIST.md")).read()
        self.assertIn("❌ FAIL", chk)
        for iid in ("U1", "U2", "U3", "U4"):
            self.assertIn(f"| {iid} |", chk)

    def test_scope_excludes_power_hardware(self):
        refs = {f["ref"] for f in self.b.footprints}
        self.assertFalse(any(r.startswith(("F", "SW", "K", "RS", "SH")) for r in refs))


if __name__ == "__main__":
    unittest.main()
