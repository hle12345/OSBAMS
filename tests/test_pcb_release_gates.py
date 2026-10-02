"""Release-gate guard for the controller PCB: no fabrication package while any gate is open; plan/bench docs present."""
import json, os, re, subprocess, sys, unittest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)
PCB = os.path.join(ROOT, "docs", "rev2", "pcb")
MFG = os.path.join(ROOT, "manufacturing")
FAB_EXT = (".gbr", ".drl", ".xlsx", ".zip", ".gtl", ".gbl")

from tools.mfg.pcb_model import Board
from tools.mfg.schem_model import Schematic
from tools.mfg import review_checks


def gates():
    return json.load(open(os.path.join(PCB, "RELEASE_GATES.json")))


class TestReleaseGates(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.b, cls.s = Board(), Schematic()

    def test_no_fabrication_files_while_a_gate_is_open(self):
        g = {k: v for k, v in gates().items() if not k.startswith("_")}
        self.assertEqual(len(g), 6)
        if all(g.values()):
            return
        found = []
        for d, _, fs in os.walk(MFG):
            found += [os.path.join(d, f) for f in fs if f.lower().endswith(FAB_EXT) or "CPL" in f or "BOM" in f.upper()]
        self.assertEqual(found, [], "fabrication outputs exist while release gates are open")

    def test_builder_refuses_without_candidate_flag(self):
        if all(v for k, v in gates().items() if not k.startswith("_")):
            self.skipTest("all gates closed")
        r = subprocess.run([sys.executable, "-m", "tools.mfg.build_package"], cwd=ROOT, capture_output=True, text=True)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("rev1_physical_observations_recorded", r.stdout + r.stderr)

    def test_bench_sheet_is_one_page_with_authoritative_v1_to_v10(self):
        data = open(os.path.join(PCB, "BENCH_V1_V10_ONE_PAGE.pdf"), "rb").read()
        self.assertEqual(len(re.findall(rb"/Type\s*/Page[^s]", data)), 1)
        md = open(os.path.join(PCB, "BENCH_V1_V10.md")).read()
        for i in range(1, 11):
            self.assertRegex(md, rf"\| V{i} \|")
        self.assertNotIn("reconstruct", md.lower())

    # ------------------------------------------------------------------ Rev.2 RC1 package
    RC = os.path.join(ROOT, "Hardware", "Rev2_Controller", "OSBAMS_Rev2_RELEASE_CANDIDATE_1")
    R2 = os.path.join(ROOT, "Hardware", "Rev2_Controller")

    def _register(self):
        return json.load(open(os.path.join(self.R2, "calc", "datasheet_inputs.json")))["inputs"]

    def test_evidence_register_uses_the_three_states_and_never_fakes_local_verification(self):
        for e in self._register():
            self.assertIn(e["evidence"], ("VERIFIED_LOCAL", "USER_RELAYED_MANUFACTURER", "UNVERIFIED"), e["id"])
            self.assertTrue(e["source"], e["id"])
            if e["evidence"] == "VERIFIED_LOCAL":
                # library facts read by the build, or a manufacturer document that is committed in the repo
                self.assertTrue(re.search(r"KiCad|read by", e["source"]) or (e["doc"] and os.path.isfile(os.path.join(ROOT, e["doc"]))), e["id"])
        ids = {e["id"] for e in self._register()}
        for must in ("vo610a_ctr_min_1mA", "vo610a_ctr_min_10mA", "ina228_pinmap", "lmr_pinmap", "iso_pinmap", "eb21a_drawing", "q_rds"):
            self.assertIn(must, ids)

    def test_rc1_package_files_exist(self):
        need = ["OSBAMS_Rev2_RC1.kicad_pro", "OSBAMS_Rev2_RC1.kicad_sch", "OSBAMS_Rev2_RC1.kicad_pcb", "OSBAMS_Rev2_RC1_Schematic.pdf",
                "OSBAMS_Rev2_RC1_BOM.xlsx", "OSBAMS_Rev2_RC1_CPL.csv", "OSBAMS_Rev2_RC1_Assembly_Drawing_Top.pdf",
                "ASSEMBLY_NOTES.md", "FABRICATION_NOTES.md", "ISOLATION_CHECK.txt", "OSBAMS_Rev2_RC1.kicad_dru", "PRE_PCBWAY_RELEASE_CHECKLIST.md", "EVIDENCE_RISK_CLASSIFICATION.md", "CONNECTOR_FOOTPRINT_CHECK.md", "OSBAMS_Rev2_RC1_Test_Point_Map.pdf", "POWER_TREE_AND_CALCULATIONS.md", "DFM_DFA_REPORT.md",
                "EVIDENCE_REGISTER.md", "SUPPLY_CHAIN_REPORT.md", "ERC_REPORT.rpt", "DRC_REPORT.rpt", "PCBWAY_RELEASE_CANDIDATE_REPORT.md", "NETLIST_CHECK.txt", "CONTROLLER_PROTECTION_AND_CONNECTOR_AUDIT.md"]
        for f in need:
            self.assertTrue(os.path.getsize(os.path.join(self.RC, f)) > 0, f)
        # the script-generated candidate Gerbers/drills must not ship: the owner exports them from the final PCB in KiCad 10
        self.assertFalse(os.path.exists(os.path.join(self.RC, "gerber")) or os.path.exists(os.path.join(self.RC, "drill")))
        self.assertEqual([f for f in os.listdir(self.RC) if f.lower().endswith((".gbr", ".drl", ".gtl", ".gbl")) or "gerber" in f.lower()], [])

    def test_erc_clean_and_drc_has_no_unrouted_nets(self):
        erc = open(os.path.join(self.RC, "ERC_REPORT.rpt")).read()
        self.assertEqual(len(re.findall(r"^\[", erc, re.M)), 0, "ERC violations present")
        drc = open(os.path.join(self.RC, "DRC_REPORT.rpt")).read()
        self.assertIn("** Found 0 unconnected pads **", drc)
        self.assertIn("** Found 0 Footprint errors **", drc)
        self.assertIn("** Found 0 DRC violations **", drc)
        dru = open(os.path.join(self.RC, "OSBAMS_Rev2_RC1.kicad_dru")).read()
        self.assertIn("isolation_host_to_controller", dru)

    def test_protection_audit_sections_and_rc12e_changes(self):
        t = open(os.path.join(self.RC, "CONTROLLER_PROTECTION_AND_CONNECTOR_AUDIT.md")).read()
        for h in ("## PASS", "## CHANGE REQUIRED", "## FIRST-ARTICLE VALIDATION", "## REMAINING FABRICATION BLOCKERS"):
            self.assertIn(h, t)
        self.assertIn("NOT RELEASED FOR FABRICATION", t)
        self.assertNotIn("PRODUCTION READY", t.upper().replace("NOT PRODUCTION READY", ""))
        # J9 uses the Samtec land pattern (0.74 x 2.79 mm pads) and the differential clamp D15 exists
        fp = open(os.path.join(self.RC, "OSBAMS_Rev2.pretty", "FTSH-105-01-L-DV-K.kicad_mod")).read()
        self.assertIn("(size 2.79 0.74)", fp)
        self.assertNotIn("(size 2.4 0.74)", fp)
        pcb = open(os.path.join(self.RC, "OSBAMS_Rev2_RC1.kicad_pcb")).read()
        self.assertIn('"OSBAMS_Rev2:FTSH-105-01-L-DV-K"', pcb)
        self.assertIn('"D15"', pcb)
        # resistor pulse capability is not published by Panasonic: it must never be marked verified
        reg = {e["id"]: e for e in self._register()}
        self.assertEqual(reg["rs_pulse_rating"]["evidence"], "UNVERIFIED")
        self.assertNotEqual(reg["rs_pulse_rating"]["evidence"], "VERIFIED_LOCAL")
        self.assertEqual(reg["ina228_diff_max"]["value"], "+/-40 V")

    def test_netlist_and_polarity_checks_pass(self):
        t = open(os.path.join(self.RC, "NETLIST_CHECK.txt")).read()
        self.assertIn("mismatches: 0", t)
        self.assertIn("polarity errors: none", t)

    def test_rc1_report_has_the_two_required_sections_and_no_release_claims(self):
        t = open(os.path.join(self.RC, "PCBWAY_RELEASE_CANDIDATE_REPORT.md")).read()
        self.assertIn("## READY FOR REVIEW", t)
        self.assertIn("## BLOCKERS BEFORE PCBWAY ORDER", t)
        for d in (self.RC, self.R2):
            for f in os.listdir(d):
                if f.endswith(".md") and f not in ("BUILD_ENVIRONMENT.md",):
                    for ln in open(os.path.join(d, f)).read().split("\n"):
                        for phrase in ("PRODUCTION READY", "HARDWARE VALIDATED", "RELEASED FOR FABRICATION"):
                            if phrase in ln.upper():
                                self.assertRegex(ln, r"(?i)\bnot\b|never|does not|no ", f"{f}: {ln[:100]}")
        for f in ("ASSEMBLY_NOTES.md", "FABRICATION_NOTES.md", "DFM_DFA_REPORT.md", "SUPPLY_CHAIN_REPORT.md"):
            self.assertIn("NOT RELEASED FOR FABRICATION", open(os.path.join(self.RC, f)).read())

    def test_bom_has_required_columns_and_every_part_has_mfr_and_mpn(self):
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(self.RC, "OSBAMS_Rev2_RC1_BOM.xlsx"))["BOM"]
        hdr = [c.value for c in ws[2]]
        for col in ("Evidence Status", "Lifecycle", "Manufacturer", "MPN", "Package", "Footprint", "Primary Supplier", "Supplier SKU", "Stock Check Date", "Approved Alternate", "PCBWay Source / Consign"):
            self.assertIn(col, hdr)
        refs = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            r = dict(zip(hdr, row))
            self.assertTrue(r["Manufacturer"] and r["MPN"], r["Reference(s)"])
            self.assertIn(r["Evidence Status"], ("VERIFIED_LOCAL", "USER_RELAYED_MANUFACTURER", "UNVERIFIED"))
            self.assertEqual(r["Stock Check Date"], "NOT CHECKED")
            refs += r["Reference(s)"].split(", ")
        sys.path.insert(0, self.R2)
        from design import rev2_design as D
        parts = {r for r, c in D.COMPS.items() if c["key"] not in ("TP", "MH", "FID")}
        self.assertEqual(set(refs), parts)

    def test_design_follows_the_locked_decisions(self):
        sys.path.insert(0, self.R2)
        from design import rev2_design as D
        nets = D.all_nets()
        self.assertNotIn("+5V", nets)
        self.assertNotIn("5V", nets)
        u1 = D.COMPS["U1"]["nets"]
        for pin, net in (("PA0", "ESTOP_SENSE"), ("PA1", "ADC_SENSE"), ("PC8", "LOAD_EN"), ("PC9", "RELAY_FB"), ("PC10", "ARM_SENSE"), ("PB8", "SCL1"), ("PB9", "SDA1"), ("PB10", "SCL2"), ("PB11", "SDA2"),
                         ("PA2", "UART_TX_MCU"), ("PA3", "UART_RX_MCU")):
            self.assertEqual(u1[pin], net)
        # the MCU reaches the coil only through R21 -> Q1 gate; sensing nets never reach the coil path
        self.assertEqual({r for r, _ in nets["LOAD_EN"]}, {"U1", "R21", "TP26"})
        for n in ("ESTOP_SENSE", "ARM_SENSE", "RELAY_FB"):
            self.assertFalse({r for r, _ in nets[n]} & {"Q1", "R21", "R22", "J2", "J3", "J4"}, n)
        # E-stop and ARM are series elements of the coil supply
        self.assertEqual(D.COMPS["J2"]["nets"], {1: "+12V", 2: "ESTOP_OUT"})
        self.assertEqual(D.COMPS["J3"]["nets"], {1: "ESTOP_OUT", 2: "COIL_V"})
        self.assertEqual(D.COMPS["J4"]["nets"], {1: "COIL_V", 2: "COIL_SW"})
        # isolation: no part other than the ISO7721 touches both GND and GND_HOST
        both = [r for r in D.COMPS if {"GND", "GND_HOST"} <= set(D.COMPS[r]["nets"].values())]
        self.assertEqual(both, ["U7"])
        # no battery discharge connector on this PCB: no connector net carries the load path
        self.assertNotIn("RELAY_OUT", [n for r in ("J2", "J3", "J4") for n in D.COMPS[r]["nets"].values()])
        # ADC clamp cannot back-feed 3V3: D8's upper diode goes to VCLAMP, never to +3V3
        self.assertEqual(D.COMPS["D8"]["nets"], {1: "GND", 2: "VCLAMP", 3: "ADC_SENSE"})
        self.assertEqual(D.COMPS["D14"]["nets"], {1: "+3V3", 3: "VCLAMP"})
        self.assertNotIn("+3V3", D.COMPS["D8"]["nets"].values())

    def test_calculations_are_current_and_have_no_assumed_20_percent_ctr(self):
        before = open(os.path.join(self.R2, "REV2_CALCULATIONS.md")).read()
        subprocess.run([sys.executable, os.path.join(self.R2, "calc", "rev2_calcs.py")], check=True, capture_output=True)
        after = open(os.path.join(self.R2, "REV2_CALCULATIONS.md")).read()
        self.assertEqual(before, after, "REV2_CALCULATIONS.md is stale: re-run calc/rev2_calcs.py")
        self.assertIn("CTR min 13 %", after)
        self.assertNotIn("assume 20 %", after.lower())
        self.assertIn("PASS", after)
        self.assertNotIn("FAIL", after)

    def test_plan_covers_all_mandatory_items_and_is_not_applied(self):
        plan = open(os.path.join(PCB, "REV2_SCHEMATIC_PCB_PLAN.md")).read()
        for kw in ("polarity", "PA0", "PC9", "PA1", "Kelvin", "ARM", "MPN", "test point", "marking"):
            self.assertIn(kw.lower(), plan.lower(), kw)
        for f in ("RELAY_VERIFICATION.md", "PI_POWER_ARCHITECTURE.md", "MANUFACTURING_STATE.md"):
            self.assertTrue(os.path.getsize(os.path.join(PCB, f)) > 0, f)

    def test_schematic_and_pcb_agree(self):
        rv = review_checks.run(self.b, self.s)
        self.assertEqual(rv["netlist_diffs"], [])
        self.assertEqual(rv["ref_footprint_mismatch"], [])
        self.assertEqual(rv["dangling_track_ends"], [])
        self.assertEqual(rv["courtyard_overlaps"], [])

    def test_diode_polarity_mismatch_is_reproducible(self):
        """Symbol pin 1 = anode (triangle base); KiCad footprint pad 1 = cathode ('K' on silk / band at pad 1)."""
        from tools.mfg.sexp import find
        for lib in ("OSBAMS:D_H",):
            sym = self.s.libs[lib]
            tri = None
            for u in find(sym, "symbol"):
                for e in u[2:]:
                    if isinstance(e, list) and e[0] == "polyline":
                        pts = [(float(p[1]), float(p[2])) for p in find(find(e, "pts")[0], "xy")]
                        if len(pts) == 4:
                            tri = pts
            apex_x = max(p[0] for p in tri); base_x = min(p[0] for p in tri)
            pin1 = [p for u in find(sym, "symbol") for p in find(u, "pin") if find(p, "number")[0][1] == "1"][0]
            self.assertLess(float(find(pin1, "at")[0][1]), base_x)      # pin 1 on the triangle-base side = anode
            self.assertGreater(apex_x, base_x)
        for ref in ("D2", "D3"):
            fp = [f for f in self.b.footprints if f["ref"] == ref][0]
            k = [t for t in fp["texts"] if t["text"] == "K"]
            self.assertTrue(k, f"{ref} footprint has no 'K' marker")
            p1 = [p for p in fp["pads"] if p["num"] == "1"][0]
            p2 = [p for p in fp["pads"] if p["num"] == "2"][0]
            d1 = ((k[0]["x"] - p1["x"]) ** 2 + (k[0]["y"] - p1["y"]) ** 2) ** .5
            d2 = ((k[0]["x"] - p2["x"]) ** 2 + (k[0]["y"] - p2["y"]) ** 2) ** .5
            self.assertLess(d1, d2)                                       # K marker sits at pad 1

    def test_nothing_claims_production_ready(self):
        for f in os.listdir(PCB):
            if f.endswith(".md") and "SNAPSHOT" not in f:
                t = open(os.path.join(PCB, f)).read()
                self.assertNotRegex(t, r"(?i)\bis production[- ]ready\b")


if __name__ == "__main__":
    unittest.main()
